import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide layout

    # Clean, Friendly Color Palette
    COLOR_BG = RGBColor(248, 250, 252)        # Soft clean light gray (#F8FAFC)
    COLOR_PRIMARY = RGBColor(30, 58, 138)      # Trustworthy Navy Blue (#1E3A8A)
    COLOR_BLUE_ACCENT = RGBColor(37, 99, 235)  # Bright Blue (#2563EB)
    COLOR_GREEN = RGBColor(5, 150, 105)        # Clean Emerald (#059669)
    COLOR_TEXT_MAIN = RGBColor(15, 23, 42)     # Dark Slate (#0F172A)
    COLOR_TEXT_MUTED = RGBColor(71, 85, 105)   # Slate Gray (#475569)
    COLOR_CARD_BG = RGBColor(255, 255, 255)    # White
    COLOR_CARD_BORDER = RGBColor(226, 232, 240)# Soft border (#E2E8F0)
    COLOR_LIGHT_BLUE = RGBColor(239, 246, 255) # Light blue card (#EFF6FF)
    COLOR_LIGHT_GREEN = RGBColor(236, 253, 245)# Light green card (#ECFDF5)

    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG
        bg_shape.line.fill.background()
        return bg_shape

    def add_header(slide, title_text, category_text, presenter_text):
        header_box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1)
        )
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
        header_box.line.color.rgb = COLOR_CARD_BORDER
        header_box.line.width = Pt(1)

        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.12)
        tf.margin_right = Inches(0.25)
        tf.margin_bottom = Inches(0.1)

        # Small top category
        p_cat = tf.paragraphs[0]
        p_cat.text = f"{category_text.upper()}   •   PRESENTER: {presenter_text.upper()}"
        p_cat.font.name = "Calibri"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_BLUE_ACCENT

        # Big clear title
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_MAIN
        p_title.space_before = Pt(3)

    def add_card(slide, left, top, width, height, title, items, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER, icon_text=""):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.margin_right = Inches(0.25)
        tf.margin_bottom = Inches(0.2)

        p_title = tf.paragraphs[0]
        full_title = f"{icon_text}  {title}".strip() if icon_text else title
        p_title.text = full_title
        p_title.font.name = "Arial"
        p_title.font.size = Pt(15)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY
        p_title.space_after = Pt(8)

        for item in items:
            p_item = tf.add_paragraph()
            p_item.text = f"•  {item}"
            p_item.font.name = "Calibri"
            p_item.font.size = Pt(12)
            p_item.font.color.rgb = COLOR_TEXT_MUTED
            p_item.space_before = Pt(4)
            p_item.space_after = Pt(4)

        return card

    # =========================================================================
    # SLIDE 1 OF 5: TITLE & TEAM ROSTER
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    s1_hero = s1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(11.733), Inches(2.2)
    )
    s1_hero.fill.solid()
    s1_hero.fill.fore_color.rgb = COLOR_PRIMARY
    s1_hero.line.fill.background()

    tf1 = s1_hero.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.4)
    tf1.margin_top = Inches(0.25)

    p1_sub = tf1.paragraphs[0]
    p1_sub.text = "FINAL PROJECT DEFENSE & SYSTEM DEMO"
    p1_sub.font.name = "Calibri"
    p1_sub.font.size = Pt(11)
    p1_sub.font.bold = True
    p1_sub.font.color.rgb = RGBColor(147, 200, 253)

    p1_h1 = tf1.add_paragraph()
    p1_h1.text = "LaundryCare Management System"
    p1_h1.font.name = "Arial"
    p1_h1.font.size = Pt(32)
    p1_h1.font.bold = True
    p1_h1.font.color.rgb = RGBColor(255, 255, 255)
    p1_h1.space_before = Pt(4)

    p1_desc = tf1.add_paragraph()
    p1_desc.text = "Doorstep Laundry Pickup & Delivery System | P2B Sayre Highway, Panadtalan, Maramag (Near Torres Capitol College Inc.)"
    p1_desc.font.name = "Calibri"
    p1_desc.font.size = Pt(14)
    p1_desc.font.color.rgb = RGBColor(224, 231, 255)
    p1_desc.space_before = Pt(4)

    # Team Members Card
    roster_box = s1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.95), Inches(7.5), Inches(4.0)
    )
    roster_box.fill.solid()
    roster_box.fill.fore_color.rgb = COLOR_CARD_BG
    roster_box.line.color.rgb = COLOR_CARD_BORDER
    roster_box.line.width = Pt(1.5)

    rtf = roster_box.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.3)
    rtf.margin_top = Inches(0.2)

    rtf_p0 = rtf.paragraphs[0]
    rtf_p0.text = "OUR TEAM & ROLES"
    rtf_p0.font.name = "Arial"
    rtf_p0.font.size = Pt(14)
    rtf_p0.font.bold = True
    rtf_p0.font.color.rgb = COLOR_PRIMARY
    rtf_p0.space_after = Pt(10)

    members = [
        ("John Michael Bretaña", "Project Leader & Admin", "Overall system design, user permissions, and deployment"),
        ("Hazil Enoc", "Front Desk Cashier", "Customer records, order encoding, and billing"),
        ("Mark Ephraim Nicor", "Laundry Operator & Rider 3", "Washing and drying schedule, deliveries across Panadtalan Puroks 1-5"),
        ("Tristan Dave M. Plaza", "Mobile App & Rider 1", "Rider mobile view, digital signature, deliveries for Torres Capitol College Inc. & Dorms"),
        ("Marvin Oclarino", "Payments & Rider 2", "Payment recording, change calculator, deliveries along Sayre Highway"),
    ]

    for m_name, m_role, m_desc in members:
        p_m = rtf.add_paragraph()
        p_m.text = f"•  {m_name} — {m_role}"
        p_m.font.name = "Calibri"
        p_m.font.size = Pt(12)
        p_m.font.bold = True
        p_m.font.color.rgb = COLOR_TEXT_MAIN

        p_sub = rtf.add_paragraph()
        p_sub.text = f"    Role: {m_desc}"
        p_sub.font.name = "Calibri"
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = COLOR_TEXT_MUTED
        p_sub.space_after = Pt(4)

    # Right Quick Summary Card
    right_box = s1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.55), Inches(2.95), Inches(3.98), Inches(4.0)
    )
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = COLOR_LIGHT_BLUE
    right_box.line.color.rgb = RGBColor(191, 219, 254)
    right_box.line.width = Pt(1.5)

    qtf = right_box.text_frame
    qtf.word_wrap = True
    qtf.margin_left = Inches(0.25)
    qtf.margin_top = Inches(0.25)

    qp0 = qtf.paragraphs[0]
    qp0.text = "QUICK PROJECT OVERVIEW"
    qp0.font.name = "Arial"
    qp0.font.size = Pt(14)
    qp0.font.bold = True
    qp0.font.color.rgb = COLOR_PRIMARY
    qp0.space_after = Pt(10)

    highlights = [
        ("📍 Location Covered", "Near Torres Capitol College Inc., P2B Sayre Highway, Panadtalan, Maramag."),
        ("🛵 3 Bajaj Delivery Motorcycles", "3 dedicated riders assigned to specific neighborhood zones."),
        ("📞 Contact & Support", "Hotline: 0906-188-8611 • Hours: 8:00 AM - 4:30 PM (Mon-Sat)."),
        ("✅ Fully Tested & Working", "Passed all 79 automated software tests with 0 errors.")
    ]

    for h_title, h_body in highlights:
        p_h = qtf.add_paragraph()
        p_h.text = h_title
        p_h.font.name = "Calibri"
        p_h.font.size = Pt(12)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_BLUE_ACCENT

        p_hb = qtf.add_paragraph()
        p_hb.text = h_body
        p_hb.font.name = "Calibri"
        p_hb.font.size = Pt(11)
        p_hb.font.color.rgb = COLOR_TEXT_MUTED
        p_hb.space_after = Pt(6)

    # =========================================================================
    # SLIDE 2 OF 5: THE PROBLEM
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Problem: What Traditional Laundry Shops Struggle With", "Why We Built This", "Hazil Enoc (Front Desk Cashier)")

    w_col = Inches(2.78)
    h_col = Inches(5.1)
    top_col = Inches(1.75)
    gap = Inches(0.2)

    c1_x = Inches(0.8)
    c2_x = c1_x + w_col + gap
    c3_x = c2_x + w_col + gap
    c4_x = c3_x + w_col + gap

    add_card(s2, c1_x, top_col, w_col, h_col, "Lost Paper Slips", [
        "Paper receipts get wet, torn, or misplaced on the counter.",
        "Handwritten tags fall off bags, causing mixed-up clothes.",
        "Hard to look up past customer records when using paper notebooks."
    ], icon_text="🏷️")

    add_card(s2, c2_x, top_col, w_col, h_col, "No Status Tracking", [
        "Customers don't know if their clothes are currently being washed or dried.",
        "Students and busy families keep calling to ask if their laundry is ready.",
        "No notifications sent when riders are already on the way."
    ], icon_text="👀")

    add_card(s2, c3_x, top_col, w_col, h_col, "Messy Delivery Dispatch", [
        "Riders only get text messages without clear addresses or landmarks.",
        "Riders overlap in the same area, wasting time and motorcycle fuel.",
        "Too many pickup requests during peak hours can overwhelm riders."
    ], icon_text="🛵")

    add_card(s2, c4_x, top_col, w_col, h_col, "Manual Math & Change Errors", [
        "Riders calculate rates and change in their head at the doorstep.",
        "Mental math mistakes cause short cash at the end of the day.",
        "No digital signature or proof that the laundry was actually received."
    ], icon_text="🧮")

    # =========================================================================
    # SLIDE 3 OF 5: THE SOLUTION (3 SIMPLE PORTALS)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "The Solution: 3 Simple Portals for Everyone", "Our Solution", "Mark Ephraim Nicor (Laundry Operator & Rider 3)")

    w_card3 = Inches(3.77)
    h_card3 = Inches(5.1)
    gap3 = Inches(0.2)

    p1_x = Inches(0.8)
    p2_x = p1_x + w_card3 + gap3
    p3_x = p2_x + w_card3 + gap3

    add_card(s3, p1_x, top_col, w_card3, h_card3, "1. For Customers", [
        "No Account Needed: Anyone can book without complicated registration.",
        "Book in 1 Minute: Choose your neighborhood zone, date, and preferred time.",
        "Live Laundry Tracker: Enter your tracking code to see 5 real-time stages from intake to doorstep delivery.",
        "Easy Payment Info: Clear rates (₱50/kilo) with Cash on Delivery or GCash options."
    ], icon_text="📱")

    add_card(s3, p2_x, top_col, w_card3, h_card3, "2. For Shop Staff", [
        "Customer Records: Easily search customer phone numbers and past orders.",
        "Order Encoding: Staff enters the weight in kg, and the system auto-computes the total bill.",
        "Machine Telemetry: Shows how many washing machines and dryers are currently busy or free.",
        "Cashier Payments: Record customer payments (Cash/GCash) and print paper receipts."
    ], icon_text="🧺")

    add_card(s3, p3_x, top_col, w_card3, h_card3, "3. For Delivery Riders", [
        "Mobile-Friendly Screen: Designed to be easily used while mounted on a motorcycle.",
        "1-Tap GPS Button: Opens the customer's location directly in Google Maps or Waze.",
        "Call Customer: One tap to call the customer's phone when arriving.",
        "Screen Signature: Customer signs with their finger on the screen as proof of delivery."
    ], icon_text="🛵")

    # =========================================================================
    # SLIDE 4 OF 5: HOW IT WORKS & SECURITY
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "How We Built & Tested the System", "Engineering & Testing", "Tristan Dave M. Plaza (Mobile App & Rider 1)")

    # Left: Simple Architecture
    add_card(s4, Inches(0.8), top_col, Inches(5.76), Inches(5.1), "How the System Works Behind the Scenes", [
        "Built with Python & Flask: A fast and lightweight web framework that runs smoothly on laptops and phones.",
        "Organized Code: The project is divided into clean modules (Customers, Orders, Pickups, Deliveries, Payments) so team members can work together easily.",
        "Secure Role Logins: Admin, Staff, and Riders each have their own login so staff only see what they need.",
        "Reliable Database: Uses SQLite to store all orders, payments, and customer records safely."
    ], icon_text="💻")

    # Right: Testing & Safety
    add_card(s4, Inches(6.76), top_col, Inches(5.77), Inches(5.1), "System Safety & Automated Testing", [
        "No Invalid Inputs: System blocks accidental mistakes like negative weights, past dates, or empty names.",
        "79 / 79 Automated Tests Passed: We created 79 test programs to test every button, route, and calculation with zero errors.",
        "Safe and Private: Protects against web attacks and prevents sensitive business reports from leaking.",
        "Always Online: Deployed live on the cloud (Render), ready to use anytime from anywhere across Panadtalan and school campus."
    ], icon_text="🛡️")

    # =========================================================================
    # SLIDE 5 OF 5: MOTORCYCLE LOGISTICS & LIVE DEMO
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Motorcycle Logistics & Live System Demonstration", "Operations & Demo", "Marvin Oclarino (Payments & Rider 2)")

    # Left: Bajaj Motorcycle Logistics
    add_card(s5, Inches(0.8), top_col, Inches(5.76), Inches(3.6), "Bajaj Motorcycle Delivery in Panadtalan", [
        "Central Hub: Main laundry hub located at P2B Sayre Highway, Panadtalan (Near Torres Capitol College Inc.).",
        "Zone 1 (Torres Capitol College Inc. & Dorms): Handled by Tristan Dave Plaza (Bajaj Moto 1).",
        "Zone 2 (Sayre Highway Commercial Corridor): Handled by Marvin Oclarino (Bajaj Moto 2).",
        "Zone 3 (Panadtalan Puroks 1 to 5 Residential): Handled by Mark Ephraim Nicor (Bajaj Moto 3).",
        "Automatic Assignment: The system matches student dorms, campus gates, or puroks to the assigned courier."
    ], icon_text="🛵")

    # Right: Payment & Contact
    add_card(s5, Inches(6.76), top_col, Inches(5.77), Inches(3.6), "Payments, Receipts & Support", [
        "Change Calculator: Shows rider the exact change to give the customer to avoid mistakes.",
        "Thermal Receipts: Prints clean 80mm laundry receipts with barcodes for easy pickup.",
        "Official Hotline: 0906-188-8611 (Calls & SMS).",
        "Store Hours: 8:00 AM to 4:30 PM, Monday to Saturday (Closed Sundays).",
        "Official Facebook Page: facebook.com/LaundryCare."
    ], icon_text="💳")

    # Bottom Demo Transition Banner
    demo_banner = s5.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.6), Inches(11.733), Inches(1.3)
    )
    demo_banner.fill.solid()
    demo_banner.fill.fore_color.rgb = COLOR_LIGHT_GREEN
    demo_banner.line.color.rgb = RGBColor(167, 243, 208)
    demo_banner.line.width = Pt(1.5)

    dtf = demo_banner.text_frame
    dtf.word_wrap = True
    dtf.margin_left = Inches(0.3)
    dtf.margin_top = Inches(0.18)

    dp0 = dtf.paragraphs[0]
    dp0.text = "NOW STARTING THE LIVE SYSTEM DEMONSTRATION"
    dp0.font.name = "Arial"
    dp0.font.size = Pt(13)
    dp0.font.bold = True
    dp0.font.color.rgb = COLOR_GREEN

    dp1 = dtf.add_paragraph()
    dp1.text = "Demo Flow: 1. Customer Books Online  ➔  2. Staff Encodes Weight & Computes Bill  ➔  3. Machine Telemetry Status  ➔  4. Customer Checks Live Tracker  ➔  5. Rider Delivers & Gets Finger Signature  ➔  6. Cashier Records Payment & Prints Receipt."
    dp1.font.name = "Calibri"
    dp1.font.size = Pt(11)
    dp1.font.color.rgb = COLOR_TEXT_MAIN
    dp1.space_before = Pt(3)

    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out_dir = r"C:\Users\Saidimar Rey\.gemini\antigravity\scratch\Laundry-Pickup-and-Delivery-Management-System"
    out_file = os.path.join(out_dir, "LaundryCare_Presentation.pptx")
    create_deck(out_file)
